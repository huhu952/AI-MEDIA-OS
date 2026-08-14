# -*- coding: utf-8 -*-
"""第一阶段验收：图片 / PDF / Word / Excel / pandas。"""

import csv
import os
import sys

TMP = r"C:\Users\administered\Documents\Codex\AI-MEDIA-OS\temp\acceptance"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def main():
    os.makedirs(TMP, exist_ok=True)
    results = []

    # 1. 图片读取（Pillow）
    try:
        from PIL import Image, ImageDraw
        img = Image.new("RGB", (320, 180), "white")
        ImageDraw.Draw(img).rectangle([10, 10, 100, 60], fill="red")
        p = os.path.join(TMP, "test.png")
        img.save(p)
        im = Image.open(p)
        assert im.size == (320, 180) and im.mode == "RGB"
        results.append(("PASS", "图片读取 (Pillow)", f"{im.format} {im.size}"))
    except Exception as e:
        results.append(("FAIL", "图片读取 (Pillow)", str(e)))

    # 2. PDF 读取（PyMuPDF）
    try:
        import pymupdf
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 72), "AI-MEDIA-OS PDF test 中文测试", fontname="china-s")
        p = os.path.join(TMP, "test.pdf")
        doc.save(p)
        doc.close()
        d2 = pymupdf.open(p)
        text = d2[0].get_text()
        assert "PDF test" in text and "中文" in text
        results.append(("PASS", "PDF 读取 (PyMuPDF)", f"{len(text)} 字符"))
        d2.close()
    except Exception as e:
        results.append(("FAIL", "PDF 读取 (PyMuPDF)", str(e)))

    # 3. Word（python-docx）
    try:
        import docx
        d = docx.Document()
        d.add_heading("AI-MEDIA-OS Word 测试", 0)
        d.add_paragraph("这是 python-docx 写入的段落，用于验收测试。")
        p = os.path.join(TMP, "test.docx")
        d.save(p)
        d2 = docx.Document(p)
        text = "\n".join(par.text for par in d2.paragraphs)
        assert "AI-MEDIA-OS" in text and "验收" in text
        results.append(("PASS", "Word (python-docx)", f"{len(text)} 字符"))
    except Exception as e:
        results.append(("FAIL", "Word (python-docx)", str(e)))

    # 4. Excel（openpyxl）
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["平台", "播放", "点赞"])
        ws.append(["抖音", 12000, 350])
        ws.append(["小红书", 8000, 280])
        p = os.path.join(TMP, "test.xlsx")
        wb.save(p)
        wb2 = openpyxl.load_workbook(p)
        ws2 = wb2.active
        rows = list(ws2.iter_rows(values_only=True))
        assert rows[1][1] == 12000 and rows[2][2] == 280
        results.append(("PASS", "Excel (openpyxl)", f"{len(rows)} 行"))
    except Exception as e:
        results.append(("FAIL", "Excel (openpyxl)", str(e)))

    # 5. pandas 数据分析
    try:
        import pandas as pd
        df = pd.DataFrame({"平台": ["抖音", "抖音", "小红书"],
                           "播放": [10000, 14000, 8000]})
        p = os.path.join(TMP, "test.csv")
        df.to_csv(p, index=False, encoding="utf-8-sig")
        df2 = pd.read_csv(p)
        total = df2["播放"].sum()
        assert total == 32000
        results.append(("PASS", "pandas 数据分析", f"播放合计 {total}"))
    except Exception as e:
        results.append(("FAIL", "pandas 数据分析", str(e)))

    for status, name, detail in results:
        print(f"{status} | {name} | {detail}")
    failed = any(r[0] == "FAIL" for r in results)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
